"""Conferência inteligente entre acumuladores fiscais e razão do Domínio."""

from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _moeda(valor) -> str:
    numero = float(valor or 0)
    return f"R$ {numero:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def renderizar_conferencia_fiscal(prefixo: str, empresa: str) -> None:
    """Renderiza a conferência fiscal genérica com estado isolado por empresa."""
    import streamlit as st

    chave = re.sub(r"[^a-z0-9_]+", "_", str(prefixo).lower()).strip("_")
    base = f"fiscal_{chave}"
    st.markdown("### Conferência Fiscal × Contábil")
    st.caption(
        f"Empresa: {empresa}. Envie o Resumo por Acumulador e o Razão do mesmo período. "
        "Somente acumuladores com conta preenchida são conferidos; movimentos não fiscais "
        "são separados em alertas."
    )
    col_fiscal, col_razao = st.columns(2)
    with col_fiscal:
        arquivo_acumuladores = st.file_uploader(
            "Relatório de acumuladores do Domínio", type=["xls", "xlsx"],
            key=f"{base}_acumuladores",
        )
    with col_razao:
        arquivo_razao = st.file_uploader(
            "Razão com todas as contas", type=["xls", "xlsx"], key=f"{base}_razao",
        )
    if arquivo_acumuladores is None or arquivo_razao is None:
        st.info("Envie os dois relatórios. A conferência começará automaticamente.")
        return

    import hashlib
    assinatura = hashlib.sha256(
        arquivo_acumuladores.getvalue() + arquivo_razao.getvalue()
    ).hexdigest()
    if st.session_state.get(f"{base}_assinatura") != assinatura:
        try:
            with st.spinner("Cruzando acumuladores, contas e lançamentos do razão..."):
                resultado = processar_conferencia(
                    arquivo_acumuladores.getvalue(), arquivo_acumuladores.name,
                    arquivo_razao.getvalue(), arquivo_razao.name,
                )
            st.session_state[f"{base}_resultado"] = resultado
            st.session_state[f"{base}_assinatura"] = assinatura
            st.session_state.pop(f"{base}_erro", None)
        except Exception as erro:
            st.session_state[f"{base}_erro"] = str(erro)
            st.session_state.pop(f"{base}_resultado", None)

    if st.session_state.get(f"{base}_erro"):
        st.error("Não foi possível concluir a conferência: " + st.session_state[f"{base}_erro"])
    resultado = st.session_state.get(f"{base}_resultado")
    if not isinstance(resultado, dict):
        return
    resumo = resultado.get("resumo", pd.DataFrame())
    detalhes = resultado.get("detalhes", pd.DataFrame())
    if resumo.empty:
        st.warning("Nenhuma conta pôde ser comparada.")
        return

    conferidas = int(resumo["SITUAÇÃO"].astype(str).str.startswith("CONFERE").sum())
    alertas = int(resumo["LANÇAMENTOS EXTRAS"].sum())
    revisar = int((resumo["SITUAÇÃO"] == "REVISAR").sum())
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Contas analisadas", len(resumo))
    m2.metric("Fiscal conferido", conferidas)
    m3.metric("Lançamentos em alerta", alertas)
    m4.metric("Contas para revisar", revisar)

    periodo = resultado.get("periodo_fiscal", {})
    inicio = pd.to_datetime(periodo.get("inicio"), errors="coerce")
    periodo_nome = inicio.strftime("%m%Y") if pd.notna(inicio) else "PERIODO_ANALISADO"
    nome_empresa = re.sub(r"[^A-Za-z0-9]+", "_", str(empresa)).strip("_")[:45]
    st.download_button(
        "Baixar relatório completo da conferência",
        data=gerar_relatorio_excel(resultado),
        file_name=f"{nome_empresa}_CONFERENCIA_FISCAL_{periodo_nome}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True, key=f"{base}_download",
    )
    if revisar:
        st.error(f"{revisar} conta(s) possuem diferença fiscal e precisam de revisão.")
    elif alertas:
        st.warning("Os valores fiscais conferem, mas existem lançamentos contábeis adicionais para revisar.")
    else:
        st.success("Todas as contas conferem e não foram encontrados lançamentos adicionais.")

    aba_visao, aba_alertas, aba_todos = st.tabs([
        "Visão geral", f"Alertas ({alertas})", "Todos os lançamentos"
    ])
    with aba_visao:
        ordem = {"REVISAR": 0, "AUSENTE NO CONTÁBIL": 1, "CONFERE COM ALERTAS": 2, "CONFERE": 3}
        resumo_ordenado = resumo.assign(
            _ORDEM=resumo["SITUAÇÃO"].map(ordem).fillna(9)
        ).sort_values(["_ORDEM", "CONTA"])
        for _, item in resumo_ordenado.iterrows():
            conta, situacao = str(item["CONTA"]), str(item["SITUAÇÃO"])
            extras = int(item["LANÇAMENTOS EXTRAS"])
            titulo = f"Conta {conta} · {situacao}" + (f" · {extras} alerta(s)" if extras else "")
            with st.expander(titulo, expanded=situacao != "CONFERE"):
                c1, c2, c3, c4 = st.columns(4)
                c1.metric("Valor fiscal", _moeda(item["VALOR FISCAL"]))
                c2.metric("Contábil compatível", _moeda(item["CONTÁBIL COMPATÍVEL"]))
                c3.metric("Diferença fiscal", _moeda(item["DIFERENÇA FISCAL"]))
                c4.metric("Total movimentado", _moeda(item["TOTAL DA CONTA"]))
                movimentos = detalhes[detalhes["CONTA"].astype(str).eq(conta)].copy()
                if not movimentos.empty:
                    movimentos["DATA"] = pd.to_datetime(movimentos["DATA"]).dt.strftime("%d/%m/%Y")
                    st.dataframe(movimentos[["DATA", "HISTÓRICO", "CONTRAPARTIDA", "VALOR", "CLASSIFICAÇÃO"]], use_container_width=True, hide_index=True)
    with aba_alertas:
        quadro = detalhes[detalhes["CLASSIFICAÇÃO"].eq("ALERTA - NÃO FISCAL")].copy()
        if quadro.empty:
            st.success("Nenhum lançamento adicional foi encontrado nas contas conferidas.")
        else:
            quadro["DATA"] = pd.to_datetime(quadro["DATA"]).dt.strftime("%d/%m/%Y")
            st.dataframe(quadro[["CONTA", "DATA", "HISTÓRICO", "CONTRAPARTIDA", "VALOR"]], use_container_width=True, hide_index=True)
    with aba_todos:
        quadro = detalhes.copy()
        if not quadro.empty:
            quadro["DATA"] = pd.to_datetime(quadro["DATA"]).dt.strftime("%d/%m/%Y")
            st.dataframe(quadro, use_container_width=True, hide_index=True)


def _texto(valor) -> str:
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return ""
    return re.sub(r"\s+", " ", str(valor)).strip()


def _numero(valor) -> float:
    if isinstance(valor, (int, float)) and not isinstance(valor, bool) and not pd.isna(valor):
        return round(float(valor), 2)
    texto = _texto(valor).replace("R$", "").replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return round(float(texto), 2)
    except ValueError:
        return 0.0


def _rotulo(valor) -> str:
    texto = unicodedata.normalize("NFKD", _texto(valor))
    texto = "".join(letra for letra in texto if not unicodedata.combining(letra))
    return re.sub(r"[^A-Z0-9]+", " ", texto.upper()).strip()


def _primeiro_preenchido(valores: list, inicio: int | None, largura: int = 4):
    if inicio is None:
        return ""
    for indice in range(inicio, min(inicio + largura, len(valores))):
        if _texto(valores[indice]):
            return valores[indice]
    return ""


def _excel(conteudo: bytes, nome: str) -> pd.ExcelFile:
    try:
        return pd.ExcelFile(io.BytesIO(conteudo))
    except Exception as erro_original:
        if Path(nome).suffix.lower() != ".xls":
            raise ValueError(f"Não foi possível abrir {nome}: {erro_original}") from erro_original
        conversor = shutil.which("soffice") or shutil.which("libreoffice")
        if not conversor:
            raise ValueError(
                "O arquivo XLS veio com a estrutura interna danificada. Exporte novamente "
                "pelo Domínio em XLSX ou instale o LibreOffice no servidor."
            ) from erro_original
        pasta = tempfile.mkdtemp(prefix="razync_fiscal_")
        origem = Path(pasta) / "origem.xls"
        origem.write_bytes(conteudo)
        try:
            processo = subprocess.run(
                [conversor, "--headless", "--convert-to", "xlsx", "--outdir", pasta, str(origem)],
                capture_output=True, text=True, timeout=45, check=False,
                env={**os.environ, "HOME": pasta},
            )
            convertido = Path(pasta) / "origem.xlsx"
            if processo.returncode != 0 or not convertido.exists():
                raise ValueError("Não foi possível recuperar o XLS exportado pelo Domínio.")
            dados = convertido.read_bytes()
            return pd.ExcelFile(io.BytesIO(dados))
        finally:
            shutil.rmtree(pasta, ignore_errors=True)


def ler_acumuladores(conteudo: bytes, nome: str) -> tuple[pd.DataFrame, dict]:
    xls = _excel(conteudo, nome)
    periodo = {"inicio": None, "fim": None}
    registros = []
    for aba in xls.sheet_names:
        bruto = pd.read_excel(xls, sheet_name=aba, header=None, dtype=object)
        tipo = ""
        col_codigo, col_descricao = 0, None
        col_valor, col_conta = None, None
        for _, linha in bruto.iterrows():
            valores = linha.tolist()
            rotulos = [_rotulo(valor) for valor in valores]
            primeiro = rotulos[0] if rotulos else ""
            if primeiro == "PERIODO":
                datas = [pd.to_datetime(v, errors="coerce") for v in valores]
                datas = [d for d in datas if pd.notna(d)]
                if datas:
                    periodo = {"inicio": min(datas), "fim": max(datas)}
            if primeiro in {"ENTRADAS", "SAIDAS", "SERVICOS"}:
                tipo = {"SAIDAS": "SAÍDAS", "SERVICOS": "SERVIÇOS"}.get(primeiro, primeiro)
                continue

            # O Domínio muda a posição das colunas conforme o relatório, a empresa
            # e a quantidade de tributos selecionados. Localizamos os cabeçalhos
            # pelo nome em vez de depender das antigas colunas fixas 13 e 49.
            cabecalho = any(r in {"COD", "CODIGO"} for r in rotulos)
            if cabecalho:
                for indice, rotulo in enumerate(rotulos):
                    if rotulo in {"COD", "CODIGO"}:
                        col_codigo = indice
                    elif rotulo == "DESCRICAO":
                        col_descricao = indice
                    elif rotulo in {"VLR CONTABIL", "VALOR CONTABIL"}:
                        col_valor = indice
                    elif rotulo in {"CONTA", "CONTA CONTABIL", "CODIGO CONTA"}:
                        col_conta = indice
                continue

            codigo = _texto(_primeiro_preenchido(valores, col_codigo))
            if not codigo.isdigit():
                continue
            valor_indice = col_valor
            if valor_indice is None:
                valor_indice = 10 if tipo == "SERVIÇOS" else 13
            valor = _numero(_primeiro_preenchido(valores, valor_indice))
            if abs(valor) < 0.005:
                continue

            conta = _texto(_primeiro_preenchido(valores, col_conta))
            if not re.fullmatch(r"\d+", conta):
                # Em algumas exportações o título "Conta" desaparece, embora o
                # código permaneça na última coluna preenchida do acumulador.
                candidatos = [
                    _texto(valor_bruto) for indice, valor_bruto in enumerate(valores)
                    if indice >= valor_indice + 20 and re.fullmatch(r"\d+", _texto(valor_bruto))
                    and _texto(valor_bruto) != "0"
                ]
                conta = candidatos[-1] if candidatos else ""
            if not re.fullmatch(r"\d+", conta):
                continue

            if col_descricao is not None and len(valores) > col_descricao:
                descricao = _texto(_primeiro_preenchido(valores, col_descricao))
            else:
                descricao = next(
                    (_texto(v) for v in valores[col_codigo + 1:valor_indice] if _texto(v)), ""
                )
            registros.append({
                "TIPO": tipo, "ACUMULADOR": codigo, "DESCRIÇÃO": descricao,
                "CONTA": conta, "VALOR_FISCAL": valor,
            })
    if not registros:
        raise ValueError("Nenhum acumulador com conta contábil preenchida foi encontrado.")
    return pd.DataFrame(registros), periodo


def ler_razao(conteudo: bytes, nome: str) -> tuple[pd.DataFrame, dict]:
    xls = _excel(conteudo, nome)
    bruto = pd.read_excel(xls, sheet_name=xls.sheet_names[0], header=None, dtype=object)
    conta = ""
    descricao_conta = ""
    registros = []
    periodo = {"inicio": None, "fim": None}
    for _, linha in bruto.iterrows():
        valores = linha.tolist()
        primeiro = _texto(valores[0] if valores else "")
        if primeiro.upper() == "PERÍODO:":
            achado = re.findall(r"\d{2}/\d{2}/\d{4}", " ".join(_texto(v) for v in valores))
            if len(achado) >= 2:
                periodo = {
                    "inicio": pd.to_datetime(achado[0], dayfirst=True),
                    "fim": pd.to_datetime(achado[1], dayfirst=True),
                }
        if primeiro.upper() == "CONTA:":
            conta = re.sub(r"\D", "", _texto(valores[1] if len(valores) > 1 else ""))
            descricao_conta = _texto(valores[5] if len(valores) > 5 else "")
            continue
        data = pd.to_datetime(valores[0] if valores else None, dayfirst=True, errors="coerce")
        if not conta or pd.isna(data):
            continue
        debito = _numero(valores[8] if len(valores) > 8 else 0)
        credito = _numero(valores[9] if len(valores) > 9 else 0)
        if abs(debito) < 0.005 and abs(credito) < 0.005:
            continue
        registros.append({
            "CONTA": conta, "DESCRIÇÃO_CONTA": descricao_conta, "DATA": data.normalize(),
            "LOTE": _texto(valores[1] if len(valores) > 1 else ""),
            "HISTÓRICO": _texto(valores[2] if len(valores) > 2 else ""),
            "CONTRAPARTIDA": _texto(valores[7] if len(valores) > 7 else ""),
            "DÉBITO": debito, "CRÉDITO": credito,
        })
    if not registros:
        raise ValueError("Nenhum lançamento foi encontrado no razão.")
    return pd.DataFrame(registros), periodo


def _parece_fiscal(historico: str) -> bool:
    texto = _texto(historico).upper()
    if re.match(r"^(PAGO|PAGAMENTO|RECEBIDO|RECEBIMENTO|BAIXA)\b", texto):
        return False
    sinais = (" CF NF ", "COMPRA DE ", "VENDA DE ", "SERVIÇOS DE TERCEIROS", "SERVICOS DE TERCEIROS", "BENEFICIAMENTO")
    return any(sinal in f" {texto} " for sinal in sinais)


def conferir_fiscal_contabil(acumuladores: pd.DataFrame, razao: pd.DataFrame):
    resumos, detalhes = [], []
    agrupado = acumuladores.groupby(["CONTA", "TIPO"], as_index=False).agg(
        VALOR_FISCAL=("VALOR_FISCAL", "sum"),
        ACUMULADORES=("ACUMULADOR", lambda x: ", ".join(map(str, x))),
        DESCRIÇÕES=("DESCRIÇÃO", lambda x: " | ".join(map(str, x))),
    )
    for _, fiscal in agrupado.iterrows():
        conta = str(fiscal["CONTA"])
        lado = "CRÉDITO" if fiscal["TIPO"] == "SAÍDAS" else "DÉBITO"
        movimentos = razao[razao["CONTA"].astype(str).eq(conta)].copy()
        movimentos["VALOR_ANALISADO"] = pd.to_numeric(movimentos.get(lado), errors="coerce").fillna(0.0)
        movimentos = movimentos[movimentos["VALOR_ANALISADO"].abs() >= 0.005].copy()
        movimentos["COMPATÍVEL_FISCAL"] = movimentos["HISTÓRICO"].map(_parece_fiscal)
        valor_compativel = round(float(movimentos.loc[movimentos["COMPATÍVEL_FISCAL"], "VALOR_ANALISADO"].sum()), 2)
        valor_total = round(float(movimentos["VALOR_ANALISADO"].sum()), 2)
        valor_fiscal = round(float(fiscal["VALOR_FISCAL"]), 2)
        diferenca = round(valor_compativel - valor_fiscal, 2)
        extras = movimentos[~movimentos["COMPATÍVEL_FISCAL"]].copy()
        if abs(diferenca) <= 0.01:
            situacao = "CONFERE COM ALERTAS" if not extras.empty else "CONFERE"
        elif movimentos.empty:
            situacao = "AUSENTE NO CONTÁBIL"
        else:
            situacao = "REVISAR"
        resumos.append({
            "CONTA": conta, "TIPO": fiscal["TIPO"], "ACUMULADORES": fiscal["ACUMULADORES"],
            "VALOR FISCAL": valor_fiscal, "CONTÁBIL COMPATÍVEL": valor_compativel,
            "TOTAL DA CONTA": valor_total, "DIFERENÇA FISCAL": diferenca,
            "LANÇAMENTOS EXTRAS": len(extras), "SITUAÇÃO": situacao,
        })
        for _, mov in movimentos.iterrows():
            detalhes.append({
                "CONTA": conta, "DATA": mov["DATA"], "LOTE": mov["LOTE"],
                "HISTÓRICO": mov["HISTÓRICO"], "CONTRAPARTIDA": mov["CONTRAPARTIDA"],
                "VALOR": mov["VALOR_ANALISADO"],
                "CLASSIFICAÇÃO": "FISCAL COMPATÍVEL" if mov["COMPATÍVEL_FISCAL"] else "ALERTA - NÃO FISCAL",
            })
    return pd.DataFrame(resumos), pd.DataFrame(detalhes)


def processar_conferencia(acumuladores_bytes: bytes, acumuladores_nome: str, razao_bytes: bytes, razao_nome: str):
    acumuladores, periodo_fiscal = ler_acumuladores(acumuladores_bytes, acumuladores_nome)
    razao, periodo_razao = ler_razao(razao_bytes, razao_nome)
    resumo, detalhes = conferir_fiscal_contabil(acumuladores, razao)
    return {
        "resumo": resumo, "detalhes": detalhes, "acumuladores": acumuladores,
        "periodo_fiscal": periodo_fiscal, "periodo_razao": periodo_razao,
    }


def gerar_relatorio_excel(resultado: dict) -> bytes:
    """Gera o relatório da conferência em abas auditáveis."""
    wb = Workbook()
    wb.remove(wb.active)
    azul = PatternFill("solid", fgColor="123047")
    azul_claro = PatternFill("solid", fgColor="DCEAF4")
    amarelo = PatternFill("solid", fgColor="FFF1CC")
    vermelho = PatternFill("solid", fgColor="FADBD8")
    verde = PatternFill("solid", fgColor="DDF2E1")

    def adicionar_aba(nome: str, quadro: pd.DataFrame):
        ws = wb.create_sheet(nome)
        if quadro is None or quadro.empty:
            ws.append(["Nenhum registro encontrado"])
            return ws
        colunas = list(quadro.columns)
        ws.append(colunas)
        for celula in ws[1]:
            celula.fill = azul
            celula.font = Font(color="FFFFFF", bold=True)
            celula.alignment = Alignment(horizontal="center")
        for registro in quadro.itertuples(index=False, name=None):
            valores = []
            for valor in registro:
                if isinstance(valor, pd.Timestamp):
                    valor = valor.to_pydatetime()
                elif pd.isna(valor):
                    valor = ""
                valores.append(valor)
            ws.append(valores)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for indice, coluna in enumerate(colunas, start=1):
            valores = [str(ws.cell(linha, indice).value or "") for linha in range(1, ws.max_row + 1)]
            ws.column_dimensions[get_column_letter(indice)].width = min(max(len(v) for v in valores) + 2, 55)
            if any(chave in str(coluna).upper() for chave in ("VALOR", "TOTAL", "DIFERENÇA", "CONTÁBIL")):
                for linha in range(2, ws.max_row + 1):
                    ws.cell(linha, indice).number_format = 'R$ #,##0.00'
            if str(coluna).upper() == "DATA":
                for linha in range(2, ws.max_row + 1):
                    ws.cell(linha, indice).number_format = "dd/mm/yyyy"
        return ws

    resumo = resultado.get("resumo", pd.DataFrame()).copy()
    detalhes = resultado.get("detalhes", pd.DataFrame()).copy()
    acumuladores = resultado.get("acumuladores", pd.DataFrame()).copy()
    ws_resumo = adicionar_aba("Resumo por conta", resumo)
    if not resumo.empty and "SITUAÇÃO" in resumo.columns:
        coluna_situacao = list(resumo.columns).index("SITUAÇÃO") + 1
        for linha in range(2, ws_resumo.max_row + 1):
            situacao = str(ws_resumo.cell(linha, coluna_situacao).value or "")
            preenchimento = (
                verde if situacao == "CONFERE" else
                amarelo if situacao == "CONFERE COM ALERTAS" else vermelho
            )
            for celula in ws_resumo[linha]:
                celula.fill = preenchimento
    alertas = detalhes[
        detalhes.get("CLASSIFICAÇÃO", pd.Series(dtype=str)).eq("ALERTA - NÃO FISCAL")
    ].copy() if not detalhes.empty else pd.DataFrame()
    ws_alertas = adicionar_aba("Alertas", alertas)
    if ws_alertas.max_row > 1:
        for linha in range(2, ws_alertas.max_row + 1):
            for celula in ws_alertas[linha]:
                celula.fill = amarelo
    adicionar_aba("Todos os lançamentos", detalhes)
    adicionar_aba("Acumuladores considerados", acumuladores)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()
